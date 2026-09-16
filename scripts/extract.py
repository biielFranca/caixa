# -*- coding: utf-8 -*-
"""Le FECHAMENTO_DE_CAIXAAA.xlsx e gera data/extract.json + data/import-report.md"""
import openpyxl, re, json, sys, unicodedata, datetime
from collections import OrderedDict

SRC = sys.argv[1] if len(sys.argv) > 1 else "FECHAMENTO_DE_CAIXAAA.xlsx"
OUT = sys.argv[2] if len(sys.argv) > 2 else "data"

def norm(s):
    if not isinstance(s, str): return ""
    s = unicodedata.normalize("NFKD", s)
    s = "".join(c for c in s if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", s).strip().lower()

def num(v):
    if isinstance(v, (int, float)) and not isinstance(v, bool): return float(v)
    if isinstance(v, str):
        t = v.strip().replace("R$", "").replace("$", "").strip()
        t = re.sub(r"[^\d,.\-]", "", t)
        if not t: return None
        if "," in t and "." in t: t = t.replace(".", "").replace(",", ".")
        elif "," in t: t = t.replace(",", ".")
        try: return float(t)
        except ValueError: return None
    return None

# ---------------------------------------------------------------- datas
def parse_daymonth(name):
    """Retorna (dia, mes, ano|None) a partir do nome da aba."""
    s = norm(name)
    m = re.search(r"(\d{1,2})\s*[-/]\s*(\d{1,2})\s*[-/]\s*(\d{4})", s)
    if m: return int(m.group(1)), int(m.group(2)), int(m.group(3))
    m = re.search(r"(\d{1,2})\s*[-/]\s*(\d{1,2})\s*[-/]\s*(\d{2})\b", s)
    if m: return int(m.group(1)), int(m.group(2)), 2000 + int(m.group(3))
    m = re.search(r"(\d{1,2})\s*[-/]\s*(\d{1,2})", s)
    if m: return int(m.group(1)), int(m.group(2)), None
    return None

BASE_YEAR = 2025

def resolve_dates(names):
    """Resolve cada aba para uma data, reparando typos, viradas de ano e abas sem nome."""
    resolved = OrderedDict()
    notes = {}
    year, prev_m, prev_d = BASE_YEAR, None, None
    unnamed_run = []

    def flush_unnamed(next_date):
        """Distribui as abas sem nome no intervalo entre a anterior e a proxima."""
        if not unnamed_run: return
        start = resolved[list(resolved)[-1]] if resolved else None
        gap = []
        if start and next_date:
            d = start + datetime.timedelta(days=1)
            while d < next_date:
                gap.append(d); d += datetime.timedelta(days=1)
        exact = len(gap) == len(unnamed_run)
        for i, sheet in enumerate(unnamed_run):
            if i < len(gap):
                resolved[sheet] = gap[i]
            else:
                last = resolved[list(resolved)[-1]]
                resolved[sheet] = last + datetime.timedelta(days=1)
            notes[sheet] = ("aba sem nome; data inferida pela posicao (intervalo exato)"
                            if exact else
                            "aba sem nome; data inferida pela posicao (intervalo AMBIGUO - revisar)")
            if not exact: notes[sheet] += "|review"
        unnamed_run.clear()

    for name in names:
        if norm(name) == "dinheiro guardado":
            continue
        p = parse_daymonth(name)
        if p is None:
            unnamed_run.append(name)
            continue

        d, mth, yr = p
        note = None
        if yr is not None:
            year = yr
        elif prev_m is not None:
            if mth == 1 and prev_m == 12:
                year += 1                                  # virada de ano legitima
            elif mth < prev_m or mth == 0:
                note = f"mes invalido/typo no nome ({mth:02d}); corrigido para {prev_m:02d}"
                mth = prev_m                               # typo de mes
        if mth == 0: mth = prev_m or 1
        try:
            date = datetime.date(year, mth, d)
        except ValueError:
            note = f"data invalida no nome ({d:02d}-{mth:02d}-{year}); usado dia seguinte ao anterior"
            last = resolved[list(resolved)[-1]]
            date = last + datetime.timedelta(days=1)

        flush_unnamed(date)
        while date in resolved.values():                   # colisao -> empurra
            note = (note or "") + "|data duplicada; deslocada"
            date += datetime.timedelta(days=1)
        resolved[name] = date
        if note: notes[name] = note
        prev_m, prev_d = date.month, date.day

    flush_unnamed(None)
    return resolved, notes

# ---------------------------------------------------------------- blocos
def find_label(ws, wanted, max_row=45, max_col=20):
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, max_row), max_col=max_col):
        for c in row:
            if norm(c.value) in wanted:
                yield c

def read_column_entries(ws, col, start_row, end_row):
    out = []
    for r in range(start_row, end_row + 1):
        v = ws.cell(row=r, column=col).value
        n = num(v)
        if n is not None and n != 0:
            note = ws.cell(row=r, column=col + 1).value
            out.append({"amount": round(n, 2), "note": (str(note).strip() if isinstance(note, str) and note.strip() else None)})
    return out


SUM_RE = re.compile(r"SUM\(\s*\$?[A-Z]{1,2}\$?(\d+)\s*:\s*\$?[A-Z]{1,2}\$?(\d+)\s*\)", re.I)

def sum_range(wsf, col, head_row):
    """Le o range da formula =SUM(X5:X18) que a planilha usa para somar a coluna."""
    if wsf is None: return None, None
    for r in range(max(1, head_row - 4), head_row + 2):
        v = wsf.cell(row=r, column=col).value
        if isinstance(v, str) and v.lstrip().startswith("="):
            m = SUM_RE.search(v)
            if m:
                lo, hi = int(m.group(1)), int(m.group(2))
                return (lo, min(hi, lo + 40)) if hi >= lo else (lo, lo)
    return None, None

def extract_day(ws, wsf):
    d = {"entries": [], "shifts": [], "platforms": [], "warnings": []}

    # abertura: rotulo na coluna A/D, valor na coluna seguinte
    for c in find_label(ws, {"din inicial", "dinheiro", "dinheiro inicial"}, max_row=3, max_col=3):
        n = num(ws.cell(row=c.row, column=c.column + 1).value)
        if n is not None: d["cash_open"] = round(n, 2); break
    for c in find_label(ws, {"moeda inicial"}, max_row=3, max_col=6):
        n = num(ws.cell(row=c.row, column=c.column + 1).value)
        if n is not None: d["coin_open"] = round(n, 2); break
    d.setdefault("cash_open", 0.0); d.setdefault("coin_open", 0.0)

    # lancamentos: cabecalhos DINHEIRO(B) / PIX(E) / CARTAO(H)
    heads = {"dinheiro": "dinheiro", "pix": "pix", "cartao": "cartao"}
    found_cols = {}
    for c in find_label(ws, set(heads), max_row=8, max_col=12):
        key = heads[norm(c.value)]
        if key not in found_cols and c.column in (2, 5, 8):
            found_cols[key] = (c.column, c.row)
    for key, (col, hrow) in found_cols.items():
        lo, hi = sum_range(wsf, col, hrow)
        if lo is None:
            lo, hi = hrow + 1, hrow + 14
            d["warnings"].append(f"coluna {key}: sem formula de soma; usado intervalo padrao")
        for e in read_column_entries(ws, col, lo, hi):
            d["entries"].append({"method": key, **e})

    # bloco CAIXA (coluna J rotulo, K valor) - apenas contagens manuais
    for c in find_label(ws, {"din", "moeda"}, max_row=12, max_col=12):
        if c.column != 10: continue
        n = num(ws.cell(row=c.row, column=c.column + 1).value)
        if n is None: continue
        if norm(c.value) == "din": d.setdefault("counted_cash", round(n, 2))
        else: d.setdefault("counted_coin", round(n, 2))
    d.setdefault("counted_cash", 0.0); d.setdefault("counted_coin", 0.0)

    # FUNCIONARIO -> TOTAL:
    start = end = None
    for c in find_label(ws, {"funcionario", "funcionarios"}, max_col=12):
        if c.column == 10: start = c.row; break
    if start:
        for r in range(start + 1, min(ws.max_row, start + 30) + 1):
            if norm(ws.cell(row=r, column=10).value) in ("total:", "total"):
                end = r; break
        end = end or start + 15
        for r in range(start + 1, end):
            raw = ws.cell(row=r, column=10).value
            name = str(raw).strip().rstrip(":").strip() if isinstance(raw, str) else None
            if not name: continue
            k = num(ws.cell(row=r, column=11).value)   # K: entregas OU valor (cozinha)
            l = num(ws.cell(row=r, column=12).value)   # L: valor calculado
            note = ws.cell(row=r, column=13).value     # M: observacao
            is_kitchen = norm(name) in ("cozinha", "cozinha:")
            d["shifts"].append({
                "name": name,
                "is_kitchen": is_kitchen,
                "deliveries": None if is_kitchen else (int(k) if k is not None and float(k).is_integer() and k < 200 else None),
                "amount": round(l, 2) if l else (round(k, 2) if is_kitchen and k else None),
                "note": str(note).strip() if isinstance(note, str) and note.strip() else None,
            })

    # TOTAL LIVRE (coluna M ou P)
    for c in find_label(ws, {"total livre"}, max_col=20):
        col = c.column
        for r in range(c.row + 1, min(ws.max_row, c.row + 8) + 1):
            lab = norm(ws.cell(row=r, column=col).value)
            if not lab or lab in ("total:", "total"): 
                if lab: break
                continue
            n = num(ws.cell(row=r, column=col + 1).value)
            if n is None: continue
            key = ("99" if "99" in lab else "ifood" if "ifood" in lab
                   else "keeta" if "keeta" in lab else "caderno" if "caderno" in lab else None)
            if key: d["platforms"].append({"platform": key, "amount": round(n, 2)})
        break
    return d

# ---------------------------------------------------------------- guardado
def extract_savings(ws):
    out, warn = [], []
    people = {}
    for c in ws[1]:
        if isinstance(c.value, str) and c.value.strip() and norm(c.value) != "coluna1":
            people[c.column] = c.value.strip()
    for col, person in people.items():
        for r in range(2, ws.max_row + 1):
            v = ws.cell(row=r, column=col).value
            if v in (None, "", " "): continue
            s = str(v).strip()
            m = re.match(r"^(\d{1,2})\s*[-/]\s*(\d{1,2})\s*(.*)$", s)
            amount = num(m.group(3)) if m else num(s)
            if amount is None:
                warn.append(f"guardado: '{person}' celula {ws.cell(row=r,column=col).coordinate} nao parseada: {s!r}")
                continue
            date = None
            if m:
                try: date = datetime.date(2026, int(m.group(2)), int(m.group(1))).isoformat()
                except ValueError: pass
            out.append({"person": person, "date": date, "amount": round(amount, 2), "note": s})
    return out, warn

# ---------------------------------------------------------------- main
def main():
    wb = openpyxl.load_workbook(SRC, data_only=True)
    wbf = openpyxl.load_workbook(SRC, data_only=False)
    dates, notes = resolve_dates(wb.sheetnames)

    days, report_rows = [], []
    for sheet, date in dates.items():
        ws = wb[sheet]
        nonempty = sum(1 for row in ws.iter_rows(max_row=min(ws.max_row, 50)) for c in row if c.value not in (None, "", " "))
        if nonempty == 0:
            report_rows.append((date.isoformat(), sheet, "IGNORADA", "aba vazia")); continue
        d = extract_day(ws, wbf[sheet])
        d["date"] = date.isoformat()
        d["source_sheet"] = sheet
        note = notes.get(sheet, "")
        d["needs_review"] = "review" in note or "AMBIGUO" in note or "duplicada" in note
        # conferencia
        caderno = round(sum(e["amount"] for e in d["entries"]), 2)
        caixa = round(d["counted_cash"] + d["counted_coin"]
                      + sum(e["amount"] for e in d["entries"] if e["method"] in ("pix", "cartao"))
                      - d["cash_open"] - d["coin_open"], 2)
        d["check_caderno"], d["check_caixa"], d["check_diff"] = caderno, caixa, round(caixa - caderno, 2)
        days.append(d)
        # needs_review = problema de IMPORTACAO. Diferenca de caixa e dado do negocio, nao erro.
        d["needs_review"] = d["needs_review"] or bool(d["warnings"])
        status = "REVISAR" if d["needs_review"] else "ok"
        report_rows.append((date.isoformat(), sheet, status, " / ".join([note.replace("|review", "")] + d["warnings"]).strip(" /")))

    savings, swarn = extract_savings(wb["dinheiro guardado"])

    import os
    os.makedirs(OUT, exist_ok=True)
    with open(f"{OUT}/extract.json", "w", encoding="utf-8") as f:
        json.dump({"days": days, "savings": savings}, f, ensure_ascii=False, indent=1)

    flagged = [r for r in report_rows if r[2] != "ok"]
    with open(f"{OUT}/import-report.md", "w", encoding="utf-8") as f:
        f.write("# Relatorio de importacao\n\n")
        f.write(f"- Abas no arquivo: **{len(wb.sheetnames)}**\n")
        f.write(f"- Dias importados: **{len(days)}**\n")
        f.write(f"- Lancamentos: **{sum(len(d['entries']) for d in days)}**\n")
        f.write(f"- Turnos de funcionario: **{sum(len(d['shifts']) for d in days)}**\n")
        f.write(f"- Registros de dinheiro guardado: **{len(savings)}**\n")
        f.write(f"- Dias marcados para revisao: **{len(flagged)}**\n")
        big = [d for d in days if abs(d["check_diff"]) > 1]
        f.write(f"- Dias com diferenca de caixa acima de R$ 1,00: **{len(big)}** "
                f"(dado real do negocio, ja presente na planilha - nao e erro de importacao)\n\n")
        f.write("## Dias que precisam de revisao\n\n| Data | Aba original | Status | Motivo |\n|---|---|---|---|\n")
        for r in flagged:
            f.write(f"| {r[0]} | `{r[1]}` | {r[2]} | {r[3] or '-'} |\n")
        if swarn:
            f.write("\n## Avisos em 'dinheiro guardado'\n\n")
            for w in swarn: f.write(f"- {w}\n")
    print(f"dias={len(days)} lancamentos={sum(len(d['entries']) for d in days)} guardado={len(savings)} revisar={len(flagged)}")

main()
